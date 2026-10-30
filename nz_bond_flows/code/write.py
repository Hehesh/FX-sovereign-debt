"""Step 7: checks + write the workbook and CSV."""
import sys
import numpy as np
import pandas as pd
sys.path.insert(0, sys.argv[1])
import load as L
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

x = pd.read_pickle(sys.argv[2])
OUTX, OUTC = sys.argv[3], sys.argv[4]
D, P, E, H, bonds, kt, hol, bdays = x["D"], x["P"], x["E"], x["H"], x["bonds"], x["kt"], x["hol"], x["bdays"]
START, END, LAG = x["START"], x["END"], x["TENDER_LAG"]

# ---------------- checks ----------------
def add_bd(d, k):
    return bdays[bdays.searchsorted(pd.Timestamp(d), side="right") + k - 1]
tot = H.pivot_table(index="month_end", columns="bond", values="total").fillna(0)
dtot = tot.diff().iloc[1:]
mats = H.groupby("bond").maturity.first()
tend, rep = L.tenders(), L.repurchases()
other = E[~E.source.str.contains("tender|Repurchase-history: Tender", case=False, regex=True)]

def recon(lag):
    ev = [(add_bd(d, lag), b, a) for d, b, a in zip(tend.trade_date, tend.bond, tend.amount)]
    ev += [(add_bd(d, lag), b, -a) for d, b, a in zip(rep.trade_date, rep.bond, rep.amount)]
    ev += list(zip(other.settle_date, other.bond, -other.face_signed))
    e = pd.DataFrame(ev, columns=["d", "b", "a"])
    e["m"] = e.d + pd.offsets.MonthEnd(0)
    em = e.groupby(["m", "b"]).a.sum().unstack().reindex(index=dtot.index, columns=dtot.columns).fillna(0)
    r = dtot - em
    for b, m in mats.items():
        mm = m + pd.offsets.MonthEnd(0)
        if mm in r.index:
            r.loc[mm, b] += tot.shift(1).loc[mm, b]
    live = ((tot.shift(1) > 0) | (tot > 0)).iloc[1:]
    rr = r.where(live).stack().dropna()
    return rr

lag_rows = []
for lag in [1, 2, 3, 4]:
    rr = recon(lag)
    lag_rows.append((f"T+{lag}", len(rr), int((rr.abs() < 0.5).sum()), int((rr.abs() >= 0.5).sum()),
                     round(rr.abs().sum(), 1)))
rr = recon(LAG)
breaks = rr[rr.abs() >= 0.5].reset_index()
breaks.columns = ["month_end", "bond", "unexplained_change_NZDm"]
why = {("GB0201",): "Pre-2009 buybacks before maturity (repurchase file starts 2009)",
       ("GB0302",): "Pre-2009 buybacks before maturity (repurchase file starts 2009)",
       ("GB0433",): "Source noise: +150/-150 alternating in NZDM month-end file",
       ("IIB0216",): "Small pre-2009 change with no matching record",
       ("GB0423",): "Small change with no matching record"}
breaks["likely_reason"] = [why.get((b,), "") for b in breaks.bond]

P["tot_face"] = P.total * P.rate_per_payment
yr = P.assign(y=P.scheduled_date.dt.year).pivot_table(index="y", columns="flow", values="tot_face", aggfunc="sum") / 1000
amar = {2000:(1.8,2.8),2001:(1.9,2.6),2002:(1.9,2.4),2003:(1.9,3.7),2004:(1.8,3.6),2005:(1.8,3.8),2006:(1.9,7.4),
        2007:(1.6,0),2008:(1.8,3.8),2009:(2.0,5.8),2010:(2.4,0),2011:(3.3,8.8),2012:(3.6,0),2013:(3.8,11.3),
        2014:(3.6,0),2015:(3.5,7.7),2016:(3.4,1.3),2017:(3.3,5.9),2018:(3.1,0),2019:(3.0,9.6),2020:(3.4,5.8),
        2021:(4.0,11.3),2022:(4.0,0),2023:(4.4,16.2),2024:(5.2,14.0),2025:(6.2,18.3),2026:(3.5,7.9)}
yr = yr.fillna(0).reset_index()
yr.columns = ["year", "coupons_all_holders_face_NZDbn", "principal_all_holders_face_NZDbn"]
yr["amar_coupons_NZDbn"] = [amar[y][0] for y in yr.year]
yr["amar_principal_NZDbn"] = [amar[y][1] for y in yr.year]
yr["coupon_diff"] = (yr.iloc[:, 1] - yr.amar_coupons_NZDbn).round(2)
yr["principal_diff"] = (yr.iloc[:, 2] - yr.amar_principal_NZDbn).round(2)
yr = yr.round(2)

# Kt formula validation on bonds where NZDM publishes Kt
cpi = L.cpi().set_index("date").cpi; cpi.index = cpi.index.to_period("Q"); chg = cpi.pct_change() * 100
kv = []
for b in ["IIB0925", "IIB0930"]:
    g = kt[kt.bond == b].reset_index(drop=True)
    k = g.kt.iloc[0]
    for i in range(1, len(g)):
        pq = g.coupon_date[i].to_period("Q") - 2
        if pq not in chg.index or (pq - 1) not in chg.index:
            break
        k = round(k * (1 + round((chg[pq] + chg[pq - 1]) / 2, 2) / 100), 2)
        kv.append((b, g.coupon_date[i].date(), g.kt[i], k, round(k - g.kt[i], 3)))
kv = pd.DataFrame(kv, columns=["bond", "coupon_date", "kt_published", "kt_rebuilt_from_cpi", "diff"])

comp = ["nom_coupon", "nom_principal", "nom_issuance", "nom_buyback", "iib_coupon", "iib_principal", "iib_issuance", "iib_buyback"]
ident = float((D[comp].sum(axis=1) - D.net_flow_private).abs().max())
weekend_flows = int(((~D.nz_business_day) & (D[comp].abs().sum(axis=1) > 0)).sum())

summary = pd.DataFrame([
    ("Month-end outstanding reconciles to flows (T+%d)" % LAG,
     f"{lag_rows[LAG-1][2]:,} of {lag_rows[LAG-1][1]:,} live bond-months explained exactly; {len(breaks)} breaks listed below (NZ$%s m in total)" % f"{breaks.unexplained_change_NZDm.abs().sum():,.0f}"),
    ("Annual coupon and principal totals vs Amar's document", f"max |diff| coupons {yr.coupon_diff.abs().max():.2f}bn, principal {yr.principal_diff.abs().max():.2f}bn (rounding)"),
    ("Kt rebuilt from CPI vs NZDM published Kt (2012-2017)", f"max |diff| {kv['diff'].abs().max():.2f} index points over {len(kv)} dates"),
    ("net_flow_private = sum of the 8 component columns", f"max |diff| {ident:.6f}"),
    ("Flows dated on a weekend or NZ holiday", f"{weekend_flows}"),
    ("Payments (coupon + principal rows)", f"{len(P):,} rows; {P.bond.nunique()} bonds; first {P.pay_date.min().date()}, last {P.pay_date.max().date()}"),
], columns=["check", "result"])

gaps = pd.DataFrame([
    ("Non-resident split", "RBNZ D31 starts Jul-2015. nonres_payments is blank before Aug-2015."),
    ("RBNZ secondary-market holdings (incl. LSAP)", "Only observable from D31 (Jul-2015 on) as NZDM market bonds minus D31 market. Assumed zero before Jul-2015. RBNZ's own LSAP purchases from investors are monetary operations, not Government flows, so they are not in the daily columns."),
    ("Buybacks before 2009", "Not published. About NZ$1.8bn of 2000-2002 pre-maturity reductions (GB0201, GB0302) have no dated record; see breaks."),
    ("IIB0216 (4.5% Feb-2016) indexation", "Not in NZDM's factors file. Kt rebuilt from CPI (method verified on later IIBs); base 100 at 15-Nov-1995 is assumed. Affects ~NZ$6m per quarter in coupons and NZ$606m at the 15-Feb-2016 maturity."),
    ("Issue and buyback prices", "Amounts are face value (IIBs: face x Kt). Cash paid differs by the bond's price, by up to ~15% for low-coupon bonds sold at high yields."),
], columns=["gap", "detail"])

# ---------------- column map ----------------
cmap = pd.DataFrame([
    ("daily_flows", "date", "Calendar day, 1-Feb-2000 to 30-Sep-2026", "-", "-"),
    ("daily_flows", "nz_business_day", "TRUE if banks open in Wellington", "holidays sheet", "Weekday and not an NZ holiday (Step 5)"),
    ("daily_flows", "nom_coupon", "Coupons on nominal bonds paid to private holders", "onissue, hd31", "Sum of payments.private_cash where type=Nominal, flow=Coupon, by pay_date (Step 4)"),
    ("daily_flows", "nom_principal", "Nominal bonds repaid at maturity to private holders", "onissue, hd31", "Sum of payments.private_cash where type=Nominal, flow=Principal (Step 4)"),
    ("daily_flows", "nom_issuance", "New nominal bonds sold by tender or syndication (negative)", "tender, syndication", "Sum of events.cash_signed (negative) where category=Issuance, type=Nominal, by settle_date (Step 6)"),
    ("daily_flows", "nom_buyback", "Nominal bonds bought back from investors by tender", "repurchase (Tender tab)", "Sum of events.cash_signed where category=Buyback, type=Nominal (Step 6)"),
    ("daily_flows", "iib_coupon", "IIB coupons to private holders, inflation-adjusted", "onissue, hd31, IIB factors", "Face coupon x Kt/100 (Steps 3-4)"),
    ("daily_flows", "iib_principal", "IIB principal at maturity, inflation-adjusted", "onissue, hd31, IIB factors", "Face x Kt/100 (Steps 3-4)"),
    ("daily_flows", "iib_issuance", "New IIBs sold (negative), face x latest Kt", "tender (IIBs tab), syndication, IIB factors", "Step 6"),
    ("daily_flows", "iib_buyback", "IIBs bought back from investors, face x latest Kt", "repurchase (Tender tab), IIB factors", "Step 6"),
    ("daily_flows", "net_flow_private", "TARGET. Net NZ$ from Government to private holders", "-", "Sum of the 8 columns above (Step 7)"),
    ("daily_flows", "iib_coupon_face", "Memo: iib_coupon before inflation adjustment", "onissue, hd31", "Step 4"),
    ("daily_flows", "iib_principal_face", "Memo: iib_principal before inflation adjustment", "onissue, hd31", "Step 4"),
    ("daily_flows", "iib_issuance_face", "Memo: iib_issuance before inflation adjustment", "tender, syndication", "Step 6"),
    ("daily_flows", "iib_buyback_face", "Memo: iib_buyback before inflation adjustment", "repurchase", "Step 6"),
    ("daily_flows", "nonres_payments", "Memo: coupons + principal paid to non-residents (blank before Aug-2015)", "hd31", "Sum of payments.nonres_cash (Step 4)"),
    ("daily_flows", "crown_payments", "Memo: coupons + principal to RBNZ, EQC, SRESL and RBNZ secondary holdings. Excluded from target", "onissue, hd31", "Sum of payments.crown_cash (Step 4)"),
    ("daily_flows", "crown_nonmarket_issuance", "Memo: bonds issued directly to / redeemed from Crown entities. Excluded", "Non market issuance", "Step 6"),
    ("daily_flows", "crown_lsap_repurchase", "Memo: Treasury buying LSAP bonds back from RBNZ. Excluded", "repurchase (LSAP tab)", "Step 6"),
    ("payments", "holdings_month_end", "Month-end used for who holds the bond", "-", "Last month-end before scheduled_date (Step 4)"),
    ("payments", "private", "Private holdings at that month-end (NZ$m face)", "hd31 'Total in the market' from Jul-2015, else onissue 'Market Bonds'", "Step 2"),
    ("payments", "rbnz_secondary", "RBNZ holdings bought in the market (LSAP etc.)", "onissue minus hd31", "Step 2"),
    ("payments", "rate_per_payment", "Coupon/2 (nominal), coupon/4 (IIB), 1 for principal", "onissue", "Step 4"),
    ("payments", "kt", "IIB index factor for the payment date", "IIB factors (IIB0216: CPI estimate)", "Step 3"),
    ("payments", "*_face / *_cash", "Holdings x rate_per_payment; cash = face x Kt/100 for IIBs", "-", "Step 4"),
    ("events", "face_signed", "Face amount, + = Government pays out, - = Government receives", "tender, syndication, repurchase, non-market", "Step 6"),
    ("events", "cash_signed", "face_signed x Kt/100 for IIBs, else face_signed", "IIB factors", "Step 6"),
], columns=["sheet", "column", "meaning", "source file", "calculation"])

about = pd.DataFrame({"NZ Government bond flows, daily": [
    "What: every coupon, maturity, new bond sale and buyback on NZ Government bonds, 1-Feb-2000 to 30-Sep-2026, netted by day.",
    "Target column: daily_flows.net_flow_private (NZ$ millions).",
    "Sign: + = Government pays NZ$ to private bondholders (coupons, maturities, buybacks). - = private investors pay NZ$ to Government (new bonds).",
    "Private = all holders except RBNZ, EQC and SRESL (and, from Jul-2015, RBNZ's secondary-market and LSAP holdings). Intra-Crown flows sit in the crown_* memo columns and are not in the target.",
    "Dates: the day cash actually moves (weekend/holiday payments moved to the next Wellington business day). The FX effect is tested on the next trading day downstream.",
    "Sheets: daily_flows (target) <- payments + events (bond-level) <- holdings, kt_factors, bonds (cleaned inputs). column_map says where every column comes from; checks shows the reconciliation.",
    "Method: see the companion method note. Built 7 Oct 2026 from NZDM files published 14 Jul-7 Oct 2026 and RBNZ D31 published 16 Sep 2026.",
]})

# ---------------- write ----------------
Pout = P[["pay_date", "scheduled_date", "bond", "type", "coupon", "maturity", "flow", "holdings_month_end", "holder_source",
          "total", "crown_direct", "rbnz_secondary", "private", "nonres", "rate_per_payment", "kt", "kt_source",
          "private_face", "private_cash", "crown_face", "crown_cash", "nonres_face", "nonres_cash"]]
Hout = H.copy()
hols = pd.DataFrame(sorted((pd.Timestamp(d), n) for d, n in hol.items() if START.year <= d.year <= END.year + 1),
                    columns=["date", "holiday"])
sheets = [("about", about), ("daily_flows", D), ("column_map", cmap), ("payments", Pout), ("events", E),
          ("holdings", Hout), ("kt_factors", kt), ("bonds", bonds), ("checks", None), ("holidays", hols)]

with pd.ExcelWriter(OUTX, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as w:
    for name, df in sheets:
        if name == "checks":
            r = 0
            for title, t in [("Summary", summary), ("Settlement lag test (tenders and buyback tenders)",
                             pd.DataFrame(lag_rows, columns=["lag", "live_bond_months", "explained", "breaks", "sum_abs_break_NZDm"])),
                             ("Reconciliation breaks at T+%d" % LAG, breaks), ("Known gaps", gaps),
                             ("Annual totals, all holders, face value vs Amar", yr), ("Kt rebuilt from CPI vs published", kv)]:
                pd.DataFrame({title: []}).to_excel(w, sheet_name="checks", startrow=r, index=False)
                t.to_excel(w, sheet_name="checks", startrow=r + 1, index=False)
                r += len(t) + 4
        else:
            df.to_excel(w, sheet_name=name, index=False)
    wb = w.book
    head = PatternFill("solid", fgColor="1F3A5F")
    for ws in wb.worksheets:
        widths = {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
            for c in row:
                if c.value is not None:
                    widths[c.column] = max(widths.get(c.column, 0), min(len(str(c.value)), 60))
        for col, wdt in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(10, wdt + 2)
        if ws.title not in ("checks", "about"):
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), head
            ws.freeze_panes = "B2"
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "#,##0.00"
    ws = wb["about"]; ws.column_dimensions["A"].width = 140
    ws["A1"].font = Font(bold=True, size=13)
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb["daily_flows"]
    for c in ws[1]:
        if c.value == "net_flow_private":
            c.fill = PatternFill("solid", fgColor="B8860B")
    ws = wb["checks"]
    for row in ws.iter_rows():
        for c in row:
            if c.column == 1 and isinstance(c.value, str) and c.value in (
                "Summary", "Known gaps", f"Reconciliation breaks at T+{LAG}") or (isinstance(c.value, str) and c.value.startswith(("Settlement lag", "Annual totals", "Kt rebuilt"))) and c.column == 1:
                c.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 110
    for name in ["column_map"]:
        ws = wb[name]
        for col, wdt in zip("ABCDE", [12, 26, 70, 40, 60]):
            ws.column_dimensions[col].width = wdt

D.to_csv(OUTC, index=False, date_format="%Y-%m-%d")
print(summary.to_string()); print(breaks.to_string()); print(pd.DataFrame(lag_rows))
